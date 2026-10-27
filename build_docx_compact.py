# -*- coding: utf-8 -*-
"""构建《微商系统课程设计报告（第一篇）》压缩版 Word 文档（≤15 页，无配图）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.environ.get(
    "REPORT_OUT",
    r"E:/JLU/oopwork/oop-work-WeChat-merchant-system/微商系统课程设计-第一篇报告.docx")

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
ASCII_FONT = "Times New Roman"
MONO_FONT = "Consolas"

doc = Document()


def _set_run(run, font=BODY_FONT, size=12, bold=False, italic=False, color=None):
    run.font.name = ASCII_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:ascii"), ASCII_FONT if font != MONO_FONT else MONO_FONT)
    if font == MONO_FONT:
        rFonts.set(qn("w:ascii"), MONO_FONT)
        rFonts.set(qn("w:hAnsi"), MONO_FONT)
    return run


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def setup_styles():
    st = doc.styles["Normal"]
    st.font.name = ASCII_FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.25

    for name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name = ASCII_FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        h.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 6)
        h.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def h3(text):
    doc.add_heading(text, level=3)


def p(text, size=12, indent=True, align=None, bold=False, italic=False, font=BODY_FONT):
    par = doc.add_paragraph()
    if indent:
        par.paragraph_format.first_line_indent = Pt(size * 2)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Pt(0)
    r = par.add_run()
    r.text = text
    _set_run(r, font=font, size=size, bold=bold, italic=italic)
    return par


def bullets(items, size=12):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 1.25
        r = par.add_run()
        r.text = it
        _set_run(r, size=size)


def numbered(items, size=12):
    for it in items:
        par = doc.add_paragraph(style="List Number")
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 1.25
        r = par.add_run()
        r.text = it
        _set_run(r, size=size)


def note(text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(4)
    r = par.add_run()
    r.text = text
    _set_run(r, size=10, italic=True, color=(0x40, 0x40, 0x40))


def code(text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.6)
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.line_spacing = 1.1
    r = par.add_run()
    r.text = text
    _set_run(r, font=MONO_FONT, size=9)
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F5F7")
    pPr.append(shd)


def table(headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run()
        r.text = htext
        _set_run(r, size=size, bold=True)
        cell.paragraphs[0].paragraph_format.space_after = Pt(1)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(cell, "E8EEF7")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            for j, line in enumerate(str(val).split("\n")):
                para = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.12
                r = para.add_run()
                r.text = line
                _set_run(r, size=size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def toc():
    par = doc.add_paragraph()
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "【请在 Word / WPS 中按 F9 更新目录】"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)


setup_styles()

# ================================================================ 封面
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
_set_run(title.add_run(), font=HEAD_FONT, size=22, bold=True).text = "微商系统课程设计报告"

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(20)
_set_run(sub.add_run(), font=HEAD_FONT, size=13).text = \
    "第一篇　需求描述 · 相关技术讨论 · 设计初步思考 · 难度目标"

table(
    ["项目名称", "基于 C/S 架构的微商（社交电商）系统"],
    [
        ["系统类型", "PC 端 C/S 桌面应用，单店经营、多用户访问"],
        ["技术路线", "Java 服务端 + Java Swing 客户端，TCP/IP 自定义协议，SQLite 数据库"],
        ["开发方式", "独立完成"],
        ["文档版本", "v1.0"],
        ["完成日期", "2026 年 9 月 14 日"],
        ["姓名 / 学号", "王彦如 / 55250408"],
        ["班级 / 指导教师", "11 班 / 【待确认：指导教师】"],
    ],
    widths=[3.6, 12.6],
    size=10.5,
)

doc.add_page_break()

# ================================================================ 目录
h1("目录")
toc()

# ================================================================ 第一章
h1("一、项目的需求描述")

h2("1.1 项目概述与系统定位")
p("本项目拟实现一个运行于 PC 平台的 C/S（客户端/服务器）架构社交电商系统。系统采用单店经营模式"
  "（只有一个商家），但支持多用户同时访问，从而引出并发处理与资源共享等实际问题。")
p("系统的业务目标不是追求界面华丽，而是完整打通一条电商交易闭环："
  "浏览商品 → 加入购物车 → 应用促销 → 结算下单 → 售后服务 → 界面风格切换。")
p("在架构上的自我约束有四条：")
numbered([
    "不使用 Web 技术栈：客户端为独立安装的桌面程序，服务端为独立进程，通过 TCP/IP 直接通信，通信协议自定义。",
    "不使用现成框架替代核心机制：分层结构、通信协议、对象与关系数据库映射、并发控制等均从原理层面自行实现，"
    "避免被框架封装遮蔽。",
    "限制业务范围：不考虑多商家入驻、真实支付渠道、移动端适配。",
    "明确“不做清单”：多商家入驻、真实支付、移动端、商品评价、收货地址等（属于纯增删改查，不产生架构价值）。",
])

h2("1.2 功能需求")

h3("1.2.1 业务逻辑")
p("（一）业务实体", indent=False, bold=True)
p("系统涉及的核心业务实体如下：")
table(
    ["实体", "对应表", "主要属性"],
    [
        ["用户", "t_user", "编号、账号、密码哈希、昵称、角色（顾客/店主）"],
        ["商品", "t_product", "编号、名称、价格、库存、描述、上架状态"],
        ["购物车条目", "t_cart_item", "用户、商品、数量"],
        ["订单", "t_order", "订单号、用户、应付金额、状态、时间"],
        ["订单明细", "t_order_item", "订单、商品名与单价快照、数量、小计"],
        ["促销活动", "t_promotion", "名称、类型、参数、启用开关"],
        ["优惠券", "t_coupon", "名称、门槛、面额、状态"],
    ],
    widths=[2.8, 3.2, 10.6],
)
note("设计要点：订单明细中的商品名与单价采用快照方式保存，避免商品价格随促销变动导致历史订单金额漂移。")

p("（二）功能模块", indent=False, bold=True)
table(
    ["模块", "功能描述"],
    [
        ["账号", "注册、登录；密码哈希加盐存储；不保存明文"],
        ["商品浏览", "商品列表（分页与搜索）、商品详情"],
        ["购物车", "加入、修改数量、删除、查看、勾选部分结算"],
        ["促销", "打折、降价、优惠券；可叠加；可指定适用商品；服务端为权威计算方"],
        ["结算", "汇总购物车 → 应用促销 → 校验并原子扣减库存 → 生成订单"],
        ["售后", "对已成交订单发起退货申请，经店主审核后完成退款"],
        ["界面风格", "提供多套界面主题，用户可切换"],
        ["店主端", "商品维护、促销配置、订单查看、售后审核"],
    ],
    widths=[3.0, 13.6],
)

p("（三）促销规则明细", indent=False, bold=True)
table(
    ["促销类型", "具体形式"],
    [
        ["打折", "统一折扣、第 2 件 9 折第 3 件 8 折、免单"],
        ["降价", "在商品原价基础上直接下调"],
        ["优惠券", "满减（满 N 元减 M 元）、抵扣（直接减 M 元）"],
    ],
    widths=[3.2, 13.4],
)
p("叠加规则：多个促销同时生效时，按固定顺序依次计算：「先打折 → 再满减 → 最后用券」。该顺序写入配置文件"
  "而非硬编码，以保证前后端一致。顺序不同会导致计算结果不同，因此必须明确约定并服务端权威执行。")

p("（四）主业务流程", indent=False, bold=True)
p("顾客登录后浏览商品 → 加入购物车 → 提交结算 → 服务端校验并扣减库存 → 生成订单 → 返回结果。"
  "若有问题，顾客可发起售后申请；店主登录后可维护商品、配置促销、审核售后。")

h3("1.2.2 UI 描述")
p("客户端为 PC 桌面图形界面，采用 MVC 分层：界面层（View）只负责展示与采集输入，控制器层（Controller）"
  "负责调度，业务规则不写在界面代码中。客户端根据账号角色动态挂载菜单：顾客看到购买相关界面，"
  "店主看到管理相关界面。")
p("核心界面清单：")
table(
    ["界面", "主要内容", "主要交互"],
    [
        ["登录 / 注册", "账号、密码输入框", "校验通过后按角色进入主界面"],
        ["商品列表", "商品表格、搜索框、分页控件", "搜索、翻页、点击进入详情"],
        ["商品详情", "完整描述、促销活动", "选择数量、加入购物车"],
        ["购物车", "条目列表、复选框、数量增减", "修改后实时刷新小计"],
        ["结算", "结算单、优惠明细逐行展示", "确认提交订单"],
        ["交易结果", "成功/失败提示、订单号、计算明细", "查看明细、发起售后"],
        ["售后申请", "订单选择、退货原因表单", "提交申请"],
        ["设置", "界面主题选择", "即时换肤"],
        ["店主端界面", "商品维护、促销配置、订单查看、售后审核", "新增、修改、上/下架、审核"],
    ],
    widths=[3.0, 8.2, 5.4],
)
p("客户端与服务端的职责划分：")
table(
    ["事项", "客户端", "服务端"],
    [
        ["商品信息展示", "负责", "提供数据"],
        ["输入的格式校验", "负责（如数量是否为正整数）", "二次校验"],
        ["购物车增删改", "发起请求、维护本地副本", "权威存储与计算"],
        ["促销金额计算", "仅做本地预演算，结果标注“预估”", "权威计算，返回逐项计算明细"],
        ["库存扣减", "不参与", "独占"],
        ["订单生成", "不参与", "独占"],
    ],
    widths=[3.6, 7.0, 6.0],
)
p("安全性约束：客户端运行在用户机器上，其计算结果可被篡改。系统规定：客户端可预估，服务端必须复算，"
  "界面预估金额一律带“预估”字样，实际金额以服务端返回为准。")

h3("1.2.3 通信")
p("（一）通信模型", indent=False, bold=True)
p("采用 TCP 长连接，在其上运行自定义应用层协议，为双向“请求-应答”模型，并额外保留服务端主动推送通道。"
  "选择长连接是为了支持服务端在特定事件发生时主动通知客户端（如订单状态变更）。")

p("（二）报文格式", indent=False, bold=True)
p("报文由定长报文头与变长数据体组成：")
table(
    ["字段", "长度", "说明"],
    [
        ["魔数", "2 字节", "固定值 0xAA55，用于快速识别非法数据"],
        ["版本号", "1 字节", "协议版本，当前 0x01"],
        ["类型", "1 字节", "1=请求，2=应答，9=服务端推送"],
        ["功能码", "4 字节", "标识具体操作"],
        ["序列号", "4 字节", "请求与应答的配对标识"],
        ["数据长度", "4 字节", "数据体字节数，解决粘包与半包问题"],
        ["数据体", "变长", "JSON 文本序列化"],
    ],
    widths=[2.4, 2.0, 12.2],
)
note("报文头固定为 16 字节。接收方先读 16 字节头部，解析数据长度，再读数据体；长度不足则继续等待，"
     "多余则留给下一条报文。")

p("（三）功能码编码规则", indent=False, bold=True)
p("4 字节整数，以 8 位十进制数表示，结构为 XXXXYYYY：前 4 位为操作码（首位 1=请求，2=应答，9=推送，"
  "后三位为功能序号）；后 4 位为参数位（如商品编号），不需要时填 0000。功能码高 4 位按模块分段，"
  "为后续新增功能预留编码空间。")

p("（四）主要功能码", indent=False, bold=True)
table(
    ["功能码", "含义", "功能码", "含义"],
    [
        ["1001", "请求商品列表", "1005", "发起售后申请"],
        ["1002", "将商品加入购物车", "1006", "领取优惠券"],
        ["1003", "查询购物车", "1007", "用户登录"],
        ["1004", "提交结算", "1008", "查询界面主题配置"],
    ],
    widths=[2.4, 6.0, 2.4, 5.8],
)

p("（五）异常传递方式", indent=False, bold=True)
p("业务异常不通过断开连接表达，而是在应答中携带错误码与错误消息（1xxx 表示请求不合法，2xxx 表示业务规则"
  "拒绝）。采用错误码加纯文本方式，避免序列化 Java 异常对象使协议绑定语言。")

h2("1.3 性能需求")

h3("1.3.1 多层结构")
p("系统采用六层结构，各层职责单一，层与层之间通过接口或协议协作。")
table(
    ["层次", "所在端", "职责", "技术载体"],
    [
        ["表现层", "客户端", "展示、交互、主题切换", "Java Swing"],
        ["客户端业务层", "客户端", "购物车副本维护、金额预演算", "Java 业务类"],
        ["通信层", "两端", "Socket 管理、协议编解码、并发连接", "TCP/IP + 自定义协议"],
        ["服务端业务层", "服务端", "交易处理、促销策略、事件发布", "Java 业务类"],
        ["持久层", "服务端", "对象与关系数据库的映射", "DAO 接口 + JDBC 实现"],
        ["数据库层", "服务端", "数据存储 + SQL 执行 + 事务", "SQLite，可切换 MySQL"],
    ],
    widths=[2.6, 1.8, 7.4, 4.8],
)
p("分层硬性约束：")
numbered([
    "跨层调用只能通过相邻层的接口或协议，禁止跨层直连（如界面代码不得出现 SQL 语句）。",
    "每层对外只暴露接口，不暴露实现类。",
    "客户端与服务端只能通过通信协议交互，不得共享文件、内存或本地方法调用。",
])

h3("1.3.2 柔性讨论")
p("柔性指系统应对变化的适应能力。分层设计把一处的变化隔离在单一层内，不向外扩散。本项目重点体现以下六个维度：")
table(
    ["变化维度", "由哪一层隔离", "变更时的影响范围"],
    [
        ["更换数据库（SQLite→MySQL）", "持久层 + 配置", "仅修改配置文件中的驱动与连接串，业务层零改动"],
        ["更换持久化实现（JDBC→MyBatis）", "持久层", "仅替换 DAO 实现类，业务层零改动"],
        ["通信方式调整（同步→异步）", "通信层", "业务层无需感知"],
        ["更换客户端语言（Java→其他）", "通信协议", "服务端零改动，新客户端遵守同一协议即可"],
        ["更换界面主题", "表现层", "业务层零改动"],
        ["新增促销玩法", "促销引擎（策略 + 反射）", "核心交易代码零改动"],
    ],
    widths=[5.0, 3.8, 7.8],
)

h3("1.3.3 扩展能力")
p("系统设计以开闭原则（对扩展开放、对修改封闭）为准则。主要扩展点如下：")
table(
    ["扩展点", "当前实现", "可扩展为", "依赖的机制"],
    [
        ["促销玩法", "打折、降价、券", "满减、秒杀、拼团、会员价", "反射 + 策略模式：新增类 + 改配置"],
        ["促销叠加", "三类玩法按顺序叠加", "任意数量玩法叠加", "装饰模式：计算链可任意延长"],
        ["商品属性", "基础商品", "多规格、多图", "表结构扩展"],
        ["通信功能码", "十余个", "按模块持续增加", "功能码高 4 位分段，预留空间"],
        ["服务端部署", "单机单进程", "多实例、集群", "事件发布订阅机制解耦"],
    ],
    widths=[2.8, 3.6, 4.4, 5.8],
)

h2("1.4 用例图")
p("参与者：")
bullets([
    "顾客（主要参与者）：完成从浏览到售后的全部购买行为。",
    "店主（次要参与者）：负责商品信息维护、促销活动配置与售后审核；订单信息只读。",
])
p("主要用例说明：")
bullets([
    "浏览商品：顾客提交分页与搜索条件，服务端查询返回商品列表。",
    "管理购物车：加入、修改数量、删除；购物车存于服务端。",
    "应用促销策略：服务端按配置的促销链计算优惠，返回逐项明细。",
    "结算下单：校验库存 → 原子扣减 → 生成订单 → 发布交易完成事件。",
    "发起售后申请：顾客对已付款订单填写原因，生成售后单等待审核。",
    "切换界面主题：选择主题，重新加载样式资源，界面即时刷新。",
    "店主端用例：维护商品信息、配置促销活动、审核售后申请、查看订单。",
])

# ================================================================ 第二章
h1("二、可能的相关技术讨论")

h2("2.1 技术全景")
p("按系统分层，本项目可能涉及的技术如下：")
table(
    ["类别", "候选技术"],
    [
        ["编程语言", "Java（服务端与客户端主力），C++（可选跨语言客户端）"],
        ["桌面界面", "Java Swing（JDK 自带，无需额外依赖，天然支持主题切换）"],
        ["网络通信", "TCP/IP、Socket、字节流读写"],
        ["协议设计", "自定义报文（魔数/版本/类型/功能码/序列号/长度/数据体）"],
        ["序列化", "JSON（可读、易调试，不绑定语言）"],
        ["并发编程", "线程池、synchronized、ReentrantLock、BlockingQueue"],
        ["数据库", "SQLite（起步，零部署）、MySQL（可切换）"],
        ["持久化", "手写 JDBC 封装（理解映射本质），后期可引入 MyBatis 演示柔性"],
        ["设计模式", "策略、装饰、观察者、工厂、单例、状态"],
        ["反射", "Java Reflection，按配置动态装配促销策略"],
        ["建模", "UML（用例图、类图、时序图、活动图、部署图）"],
        ["工程实践", "Git 版本控制、JUnit 测试、日志"],
    ],
    widths=[2.6, 14.0],
)

h2("2.2 技术熟悉度自查")
p("笔者按当前实际水平如实自评如下：")
table(
    ["技术", "我的自评", "本系统的要求"],
    [
        ["Java 基础语法", "会点（刚学完，未独立完成过项目）", "服务端与客户端主力语言"],
        ["面向对象", "会点", "全部设计的基础"],
        ["SQL 基础", "会点", "建表、增删改查"],
        ["Socket/TCP 通信", "没接触", "通信层核心"],
        ["多线程/并发", "没接触", "并发互斥难点"],
        ["设计模式", "没接触", "促销/异步核心"],
        ["反射机制", "没接触", "促销可扩展"],
        ["JDBC", "没接触", "持久层"],
        ["GUI 编程", "会点", "客户端界面"],
        ["UML 建模", "没接触", "需求/设计文档必备"],
    ],
    widths=[3.4, 6.2, 7.0],
)
p("自我评估：本项目的技术栈中，笔者真正会点的只有 Java 语法、面向对象与 SQL 基础，其余大部分是从未接触过的"
  "新技术。这一判断直接影响了后续的技术选择——所有选型都优先考虑风险最低、原理最易讲清的方案。")

h2("2.3 不熟悉技术的解决路径")
p("针对上述没接触的技术，总体解决思路是：先跑通最小例程，再结合项目迭代深化；不盲目引入框架，"
  "优先理解底层原理。重点难点的解决路径如下：")
bullets([
    "Socket/TCP 通信：先理解 TCP 连接建立与字节流读写，跑通“客户端发一句话→服务端回一句话”的例程，"
    "再改造成收发指定长度的字节块。若进度紧张，先实现定长报文加同步阻塞的最简版本保证功能交付，"
    "稳定后再升级为不定长报文。",
    "多线程与并发：掌握 Thread/Runnable、synchronized、ReentrantLock，让多个线程同时对同一个库存变量减一，"
    "观察竞态条件，再引入数据库条件更新进行修正。库存互斥优先采用数据库条件更新，"
    "把并发正确性交给数据库事务保证。",
    "设计模式：只重点学习 5 个模式：策略、装饰、观察者、工厂、单例，另加状态模式用于售后流转。"
    "每个模式整理一页笔记，包含意图、类图与一段可运行示例。模式数量不是越多越好，"
    "讲清 2 至 3 个与业务紧密贴合的模式远胜于罗列十个用不上的模式。",
    "反射机制：掌握 Class.forName、newInstance、Method.invoke 三个 API。把类名写进配置文件，"
    "程序启动时读取配置并反射实例化该类。若反射理解困难，可先用工厂模式加注册表达到同样效果。",
    "JDBC 持久化：先手写 Connection/PreparedStatement/ResultSet 完成一次“对象存成表、表读回到对象”的完整过程，"
    "理解映射本质。完全手写 JDBC 封装即可满足要求，MyBatis 仅作为后期演示持久化柔性的对照实现。",
    "桌面 GUI 编程：掌握容器、布局管理器、常用控件与事件监听。界面不追求视觉效果，做到功能可用、结构清晰、"
    "业务代码与界面代码分离即为达标。",
    "UML 建模：只掌握四类图：用例图、类图、时序图、活动图。每学一类图，就用本项目的一个真实场景画一遍。",
])

h2("2.4 选型对比与决策")
table(
    ["决策点", "候选方案", "选定理由"],
    [
        ["服务端语言", "Java/C++/C#", "Java 生态成熟，反射、并发、数据库资料最丰富"],
        ["客户端语言", "Java/C++", "Java 与服务端同语言，避免同时学习两套技术栈"],
        ["客户端界面", "Swing/JavaFX", "Swing 为 JDK 自带无需额外依赖；样式体系天然支持主题切换"],
        ["通信方式", "短连接/长连接", "长连接需要服务端主动推送，短连接无法实现"],
        ["报文长度", "定长/不定长", "不定长加长度字段，能同时解决粘包与半包，且能承载变长描述"],
        ["序列化方式", "Java 原生/JSON", "JSON 可读、易调试；Java 原生序列化会使协议绑定到 Java 语言"],
        ["同步/异步", "同步/异步/两者并存", "同步为主加推送通道：功能以同步方式交付保证进度，推送作为异步能力落地点"],
        ["服务端并发模型", "一连接一线程/线程池/NIO", "线程池足以支撑目标并发量，原理清晰易于讲解；NIO 学习成本高"],
        ["持久层", "Hibernate/MyBatis/手写 JDBC", "手写 JDBC 起步，能真正理解映射本质；后期引入 MyBatis 演示持久化可替换"],
        ["数据库", "MySQL/SQLite/两者可切", "SQLite 起步，零部署、单文件便于交付；保留 MySQL 使换数据库不影响业务层成为可验证结论"],
        ["库存互斥", "JVM 锁/数据库条件更新/乐观锁", "数据库条件更新为主：正确性由数据库事务保证，比手写锁更可靠"],
        ["消息队列", "自实现/引入中间件", "自实现，引入中间件会增加部署故障点，且封装过深不利于理解事件分发机制"],
        ["促销实现", "单一策略/策略加装饰加反射", "策略 + 装饰 + 反射：策略负责选哪些玩法，装饰负责逐层叠加，反射负责按配置装配"],
    ],
    widths=[2.6, 4.6, 9.4],
)

# ================================================================ 第三章
h1("三、关于设计的初步思考")

h2("3.1 分层与包结构")
p("本项目拟采用先行定义接口、再写实现的方式搭建骨架。服务端包结构设想如下：")
code(
    "server\n"
    "├── comm              通信层\n"
    "│   ├── protocol      报文定义与编解码\n"
    "│   ├── handler       功能码分发\n"
    "│   └── connection    连接管理与线程池\n"
    "├── service           业务层\n"
    "│   ├── OrderService / CartService / PromotionService\n"
    "│   └── event         事件发布与订阅\n"
    "├── dao               持久层（对外只暴露接口）\n"
    "│   ├── ProductDao / UserDao / OrderDao ...  <- 接口\n"
    "│   └── jdbc.JdbcProductDao ...             <- 实现\n"
    "├── model             领域对象\n"
    "├── dto               协议传输对象（与领域对象隔离）\n"
    "└── util              配置加载、日志、连接池"
)
p("客户端包结构与之呼应，划分为 ui（视图）、controller（调度）、service（客户端业务）、comm（通信）、dto。"
  "依赖只能自上而下：ui 不得依赖 dao，dao 不得依赖 ui。")

h2("3.2 领域模型与数据模型")
p("领域对象与数据库表之间存在结构性不匹配，这正是持久层存在的意义：")
table(
    ["不匹配的表现", "具体例子"],
    [
        ["继承关系无处存放", "Java 中不同促销策略是同一父类的子类，数据库只有一张表用 type 字段区分"],
        ["对象引用与外键", "对象之间用引用关联，数据库用外键关联（order_id 列）"],
        ["多对多关系", "促销与商品是多对多，对象里是 List，数据库里必须拆成三张表"],
        ["值的粒度不同", "对象中的金额值对象在数据库里退化为 DECIMAL 字段"],
    ],
    widths=[4.0, 12.6],
)

h2("3.3 促销引擎的设计")
p("这是本系统设计的核心，也是三种设计模式协同工作的集中体现。")

p("（一）三种机制的分工", indent=False, bold=True)
table(
    ["机制", "承担职责", "解决的问题"],
    [
        ["策略模式", "把每种促销玩法封装为独立的策略类，实现统一接口", "玩法可替换、可新增、互不干扰"],
        ["装饰模式", "把多个促销按顺序逐层包裹价格计算", "玩法可叠加，且叠加顺序可控制"],
        ["反射", "依据配置文件中的类名动态创建策略对象", "新增玩法无需修改任何已有代码"],
    ],
    widths=[2.4, 8.2, 6.0],
)

p("（二）装配流程", indent=False, bold=True)
numbered([
    "读取本次结算命中的促销记录，取得其玩法标识；",
    "按标识从配置中查出全限定类名，通过 Class.forName 反射实例化该策略；",
    "按 promotion.chain 的顺序依次用装饰器包裹，形成链式调用并返回给业务层。业务层只调用一次 calculate，"
    "无需关心链上有多长。",
])
p("关于「策略管组装、装饰管变换」的含义：策略决定链由哪些环节构成，装饰决定链上每一环的变换规则。"
  "两者必须分开，因为它们的变化原因不同——增加一种玩法，变的是链的构成；改变优惠的叠加方式，"
  "变的是链上某一环的变换规则。变化原因不同，就应当分属两个独立的扩展点。")

h2("3.4 通信协议的设计")
p("（一）编解码流程", indent=False, bold=True)
p("字节流到达并缓存 → 判断缓存字节数是否 >= 16 → 若不足则继续等待 → 足够则读取报文头并校验魔数 → "
  "若校验出错则丢弃并告警 → 校验正确则解析数据长度 len → 判断缓存字节数是否 >= len → 若不足则继续等待 → "
  "足够则取出数据体，进行 JSON 解析 → 按功能码分发处理 → 构造应答写回。")

p("（二）关键约束", indent=False, bold=True)
table(
    ["约束", "说明"],
    [
        ["长度计算", "以字节为单位。JSON 中的中文在 UTF-8 下占 3 字节，若按字符串长度计算会导致解析错位"],
        ["收发切分", "必须按长度字段切分，不得假设一次 read 就是一条完整报文，TCP 是字节流协议"],
        ["序列号", "用于配对长连接上的并发请求，应答通过序列号匹配，未使用序列号的实现无法支持异步应答"],
        ["心跳与超时", "双方约定空闲超时时间，超时则判定连接失效并触发重连"],
    ],
    widths=[3.2, 13.4],
)

h2("3.5 并发与互斥的设计")
p("（一）问题的本质", indent=False, bold=True)
p("多用户同时抢购同一件限量商品时，若扣减库存的实现为“先查询、再判断、后更新”三步，则在两个线程之间会出现"
  "竞态条件，导致超卖。")

p("（二）本系统的方案", indent=False, bold=True)
p("采用数据库条件更新：把判断与扣减合并为一条原子 SQL，由数据库事务机制保证原子性。")
code("UPDATE t_product SET stock = stock - ? WHERE id = ? AND stock >= ?")
p("执行后检查受影响行数：返回 1 表示扣减成功；返回 0 表示库存不足，立即回滚事务并向客户端返回库存不足。"
  "若一条 SQL 需同时扣减多件商品，则在同一事务内逐条执行，任一条影响行数为 0 即整体回滚。")

p("（三）为什么还需要 JVM 锁", indent=False, bold=True)
p("数据库条件更新已经能保证正确性，JVM 层的锁（按商品编号分段加锁）是性能优化而非正确性保证："
  "它可以拦住同一进程内对同一商品的重复请求，减少无谓的数据库往返。必须清醒认识到——JVM 锁只在单进程内有效，"
  "一旦服务端部署为多个实例，仅有 JVM 锁就会失效，而数据库条件更新依然有效。"
  "这个对比是理解本地锁与分布式锁差异的最好例子。")

p("（四）超卖复现实验", indent=False, bold=True)
p("本项目计划把上述问题做成一个可重复运行的实验，作为并发正确性的直接证据：")
numbered([
    "准备库存为 1 的商品；",
    "编写测试程序，启动 N 个线程（如 50 个）同时对同一商品发起结算；",
    "先运行无保护版本（三步式扣减），统计成功订单数与最终库存，展示成功订单数大于初始库存的超卖现象；",
    "再运行条件更新版本，展示成功订单数恰为 1、其余均返回失败；",
    "记录两次运行的完整输出作为实验数据。",
])
note("本实验的意义：它不仅是功能验证，更是一次完整的“发现缺陷 → 定位原因 → 提出方案 → 验证修复”的工程过程。")

# ================================================================ 第四章
h1("四、难度目标与理由")

h2("4.1 难度目标定位")
p("本项目将实现难度由低到高划分为四个层次，并据此为自身定位：")
table(
    ["层次", "判定标准", "本项目"],
    [
        ["一", "功能极简，无扩展设计", "不满足"],
        ["二", "攻克三个以上技术难点", "已满足"],
        ["三", "满足全部功能与性能要求", "目标层次"],
        ["四", "在满足全部要求的基础上实现明显超出预期的个人突破", "冲刺层次"],
    ],
    widths=[1.6, 10.6, 4.4],
)
p("本项目的定位是第三层次，并冲刺第四层次。具体依据：功能上覆盖账号、商品浏览、购物车、促销、结算、售后、"
  "界面风格切换、店主端八个模块；性能上落实六层结构、六维柔性、开闭原则式扩展、并发与互斥四项要求。")
p("已具备冲击第二层次以上的难点储备（实际远超三个）：")
bullets([
    "自定义应用层协议：报文设计、编解码、粘包与半包处理。",
    "多客户端并发连接：线程池管理与连接生命周期控制。",
    "库存互斥与防超卖：条件更新加事务。",
    "促销引擎：策略、装饰、反射三种机制协同。",
    "事件发布订阅与异步解耦：观察者模式加阻塞队列。",
    "对象与关系数据库的映射：手写 JDBC 完成映射并理解其本质。",
    "界面主题可替换：表现层与业务层解耦的验证。",
    "完整的软件工程过程：需求、设计、实现、测试、文档全过程。",
])

h2("4.2 支撑最高难度目标的三件事")
p("冲刺第四层次的关键在于超出要求的部分要有可展示的证据。本项目准备以三件事作为支撑：")
bullets([
    "分层的柔性可现场验证：在演示环节现场修改配置，从 SQLite 切换到 MySQL 并重启服务端，"
    "系统正常工作且业务代码零改动。",
    "促销玩法的现场扩展：在演示环节现场新增一个促销玩法，只需新增一个策略实现类、在配置中登记一行类名映射，"
    "用户端立即出现该玩法，而核心交易代码修改行数为零。",
    "并发正确性的对比演示：展示无保护版本下的超卖现象，再展示条件更新版本下的正确结果，"
    "用数据证明并发控制的必要性。",
])

h2("4.3 现实约束与取舍")
p("必须诚实说明本项目面临的现实约束：笔者目前的 Java 水平仅为语法刚学完、未独立完成过项目，多线程、"
  "Socket 通信、设计模式、反射、JDBC、界面编程、UML 建模均属首次接触，而项目周期为六周。因此，"
  "技术选型全程遵循风险最低、原理最易讲清的原则。")
table(
    ["类别", "内容"],
    [
        ["承诺完成", "六层结构、自定义协议、六大功能模块、店主端、促销引擎、并发互斥、SQLite 持久化、界面主题切换"],
        ["加分项（视进度取舍）", "心跳与断线重连、MySQL 现场切换演示、服务端监控台、异步推送通道的完整实现"],
        ["明确不做", "多商家入驻、真实支付渠道、移动端与浏览器端、商品评价、收货地址、收藏夹"],
        ["降级预案", "若进度不足，心跳与断线重连将主动降级为不实现，并在最终报告中明确注明，而不是留一个半成品"],
    ],
    widths=[3.4, 13.2],
)

h2("4.4 进度预期")
table(
    ["周次", "目标", "关键产出"],
    [
        ["第 1 周", "需求定稿与环境预热", "本报告定稿；开发环境就绪；Socket 收发与 JDBC 读写各跑通一个最小例程"],
        ["第 2 周", "通信层打通", "协议编解码、服务端线程池、客户端登录与商品列表，两端真正连通"],
        ["第 3 周", "数据层打通", "建库建表、DAO 接口与 JDBC 实现、商品详情与购物车功能"],
        ["第 4 周", "促销引擎", "策略、装饰、反射三者协同，配置驱动；结算流程与库存扣减"],
        ["第 5 周", "难点攻坚", "异步事件分发；售后状态流转；店主端核心功能完善"],
        ["第 6 周", "收尾交付", "界面主题切换；测试补齐；报告与演示材料；答辩演练"],
    ],
    widths=[2.0, 3.6, 11.0],
)

doc.save(OUT)
print("saved:", OUT)
