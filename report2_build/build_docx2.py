# -*- coding: utf-8 -*-
"""构建《微商系统课程设计报告（第二篇）》Word 文档（≤20 页）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BUILD = r"E:/JLU/oopwork/oop-work-WeChat-merchant-system/report2_build"
FIG = os.path.join(BUILD, "fig")
OUT = os.environ.get(
    "REPORT_OUT",
    r"E:/JLU/oopwork/oop-work-WeChat-merchant-system/微商系统课程设计-第二篇报告.docx")

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
ASCII_FONT = "Times New Roman"
MONO_FONT = "Consolas"

FIGMAP = {
    "图 1": ("fig01_login_c.png", 16.5),
    "图 2": ("fig02_settle_c.png", 11.5),
    "图 3": ("fig03_codec_c.png", 13.0),
}

doc = Document()


def _set_run(run, font=BODY_FONT, size=12, bold=False, italic=False, color=None):
    run.font.name = ASCII_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:ascii"), ASCII_FONT if font != MONO_FONT else MONO_FONT)
    if font == MONO_FONT:
        rFonts.set(qn("w:ascii"), MONO_FONT)
        rFonts.set(qn("w:hAnsi"), MONO_FONT)
    return run


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def setup_styles():
    st = doc.styles["Normal"]
    st.font.name = ASCII_FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.25

    for name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name = ASCII_FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        h.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 6)
        h.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def h3(text):
    doc.add_heading(text, level=3)


def p(text, size=12, indent=True, align=None, bold=False, italic=False, font=BODY_FONT):
    par = doc.add_paragraph()
    if indent:
        par.paragraph_format.first_line_indent = Pt(size * 2)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Pt(0)
    r = par.add_run()
    r.text = text
    _set_run(r, font=font, size=size, bold=bold, italic=italic)
    return par


def bullets(items, size=12):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 1.25
        r = par.add_run()
        r.text = it
        _set_run(r, size=size)


def numbered(items, size=12):
    for it in items:
        par = doc.add_paragraph(style="List Number")
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 1.25
        r = par.add_run()
        r.text = it
        _set_run(r, size=size)


def note(text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(4)
    r = par.add_run()
    r.text = text
    _set_run(r, size=10, italic=True, color=(0x40, 0x40, 0x40))


def code(text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.6)
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.line_spacing = 1.1
    r = par.add_run()
    r.text = text
    _set_run(r, font=MONO_FONT, size=9.5)
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F5F7")
    pPr.append(shd)


def caption(text):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(6)
    r = par.add_run()
    r.text = text
    _set_run(r, size=10.5, bold=True)


def figure(key, cap):
    fname, width = FIGMAP[key]
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(0)
    par.add_run().add_picture(os.path.join(FIG, fname), width=Cm(width))
    caption(cap)


def table(headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run()
        r.text = htext
        _set_run(r, size=size, bold=True)
        cell.paragraphs[0].paragraph_format.space_after = Pt(1)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(cell, "E8EEF7")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            for j, line in enumerate(str(val).split("\n")):
                para = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.12
                r = para.add_run()
                r.text = line
                _set_run(r, size=size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def toc():
    par = doc.add_paragraph()
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "【请在 Word / WPS 中按 F9 更新目录】"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)


setup_styles()

# ================================================================ 封面
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
_set_run(title.add_run(), font=HEAD_FONT, size=22, bold=True).text = "微商系统课程设计报告"

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(20)
_set_run(sub.add_run(), font=HEAD_FONT, size=13).text = \
    "第二篇　需求进一步理解 · 相关技术讨论 · 概要设计 · 其他"

table(
    ["项目名称", "基于 C/S 架构的微商（社交电商）系统"],
    [
        ["系统类型", "PC 端 C/S 桌面应用，单店经营、多用户访问"],
        ["技术路线", "Java 服务端 + Java Swing 客户端，TCP/IP 自定义协议，SQLite 数据库"],
        ["开发方式", "独立完成"],
        ["文档版本", "v1.0"],
        ["完成日期", "2026 年 9 月 15 日"],
        ["姓名 / 学号", "王彦如 / 55250408"],
        ["班级 / 指导教师", "11 班 / 【待确认：指导教师】"],
    ],
    widths=[3.6, 12.6],
    size=10.5,
)

doc.add_page_break()

# ================================================================ 目录
h1("目录")
toc()

# ================================================================ 第一章
h1("一、项目的需求的进一步理解和描述")

h2("1.1 需求理解的推进")
p("第一篇报告完成了需求的初步梳理，确立了交易闭环、六层结构与主要技术选型。本篇在既有结论之上，"
  "把各功能的业务规则细化到可执行的粒度，并明确边界条件与异常情形，为概要设计提供依据。")

h2("1.2 核心业务的规则细化")
p("（一）账号与登录", indent=False, bold=True)
p("注册时账号唯一；密码加盐哈希后存储，盐值随机生成并随用户记录保存，不存明文；新注册用户默认角色为「顾客」。"
  "登录时服务端按账号取出盐值，对输入密码做同样哈希后比对，成功则返回用户信息与角色。"
  "店主账号由建库时预置，不开放前台注册，避免权限失控。")

p("（二）商品与购物车", indent=False, bold=True)
p("商品只有店主可新增、修改、上下架；下架商品对顾客不可见、不可加购。购物车存储于服务端；"
  "同一用户对同一商品重复加购时合并数量；数量必须为正整数；结算时支持勾选部分条目。")

p("（三）促销与结算", indent=False, bold=True)
p("多个促销同时生效时按固定顺序计算：先打折 → 再满减 → 最后用券，顺序写入配置文件。"
  "促销可指定适用商品（多对多关联），结算时仅命中的促销生效。库存校验与扣减必须在服务端事务内原子完成；"
  "客户端只做金额预演算并标注「预估」，实际金额以服务端返回为准。")

p("（四）订单与售后", indent=False, bold=True)
p("订单状态流转：待付款 → 已付款 → 已完成 / 已取消；已付款订单可发起售后，进入售后中，审核通过则回补库存并退款，"
  "拒绝则恢复为已完成。「已付款」由顾客点击确认触发，属于对支付环节的模拟。")

p("（五）店主端", indent=False, bold=True)
p("店主角色承担商品维护、促销配置、售后审核、订单只读四类操作。系统只有一套客户端程序，"
  "登录后按角色动态挂载界面，不做独立的店主端程序。")

h2("1.3 边界与约束的进一步明确")
p("系统为单店经营、多用户访问；明确不做多商家入驻、真实支付、移动端与浏览器端、商品评价、收货地址、收藏夹。"
  "客户端运行在用户机器上，其计算结果可被篡改，故「客户端可预估、服务端必须复算」是贯穿全系统的安全约束。")

h2("1.4 非功能需求的细化")
bullets([
    "并发：多用户同时结算同一商品不得超卖，由数据库条件更新保证正确性。",
    "扩展：新增促销玩法不修改核心交易代码，由策略 + 反射实现。",
    "部署：客户端与服务端为两个独立进程，协议不依赖回环地址，可拆分到两台机器。",
    "界面：提供 3 套主题，切换仅影响表现层，不触及业务层。",
])

# ================================================================ 第二章
h1("二、相关技术讨论")

h2("2.1 C/S 与 B/S 结构")
p("C/S（客户端/服务器）结构中，客户端为独立安装的桌面程序，与服务端直接通信；B/S（浏览器/服务器）结构中，"
  "客户端是浏览器，通过 HTTP 访问 Web 服务器。二者的差异决定了各自的适用场景：")
table(
    ["维度", "C/S", "B/S"],
    [
        ["客户端形态", "独立桌面程序，需安装", "浏览器，免安装、跨平台"],
        ["交互与表现力", "强，可直接使用 Swing 等富控件", "受浏览器能力限制"],
        ["通信方式", "可直接使用 TCP 自定义协议", "以 HTTP 为主，实时推送需额外机制"],
        ["升级分发", "需分发安装包", "服务端集中升级"],
        ["安全性", "客户端可控、数据链路可控", "依赖浏览器与 HTTPS"],
    ],
    widths=[3.0, 6.8, 6.8],
)
p("本项目的选择是 C/S。任务要求 PC 桌面应用、自定义 TCP 协议与 Swing 界面，这正是 C/S 的典型形态；"
  "更重要的是，C/S 才能直接练到 Socket 通信与自定义协议这一底层原理，B/S 会把这些封装在 HTTP 与容器之后。")

h2("2.2 Web 服务与 JavaEE / JSP / JavaScript（B/S 技术族）")
p("JavaEE 是服务器端企业级规范（含 Servlet/JSP/EJB），JSP 是服务器端动态页面技术，JavaScript 是浏览器端脚本语言，"
  "Web 服务（SOAP/RESTful）是服务之间通过 HTTP 交互的方式。这些技术共同服务于 B/S 架构。")
p("本项目不使用它们，原因是架构形态不同：本项目客户端是 Swing 桌面程序而非浏览器，服务端通过自定义 TCP 协议"
  "而非 HTTP 通信，因此不需要 JSP/JavaScript 渲染页面，也不需要以 Web 服务对外暴露接口。"
  "这些技术属于「经过比较后排除的候选方案」，排除理由清晰——理解它们，是为了说明本项目的选型边界是明确且可论证的。")

h2("2.3 MVC 框架")
p("MVC 将程序划分为 Model（数据与业务）、View（界面）、Controller（调度）三部分，使界面与业务解耦。"
  "本项目客户端按 MVC 组织：View 只负责展示与采集输入，Controller 负责事件调度，业务规则全部放在 Service 层，"
  "不写入界面代码。服务端虽无界面，但同样遵循「业务层与表现层解耦」的思想。"
  "需要说明的是，本项目不引入 Spring MVC 等现成框架，而是从原理层面自行实现分层，"
  "以避免被框架封装遮蔽，只学会调用而学不会设计。")

h2("2.4 持久层")
p("持久层负责对象与关系数据库之间的映射，是本项目理解 ORM 本质的关键。演进路径为：先手写 JDBC"
  "（Connection/PreparedStatement/ResultSet）完成「对象存成表、表读回对象」的完整过程，"
  "再抽象出 DAO 接口隔离业务层，后期引入 MyBatis 以演示「持久层实现可替换」的柔性。")
p("数据库采用 SQLite 起步（零部署、单文件便于交付），同时保留切换 MySQL 的能力："
  "两套建表脚本分别处理自增语法差异，分页统一使用 LIMIT ? OFFSET ?，时间统一存 ISO-8601 字符串，"
  "使方言差异收敛在建表脚本中、DAO 代码不分叉。")

h2("2.5 设计模式")
p("本项目重点使用的设计模式及落点如下：")
table(
    ["模式", "在本项目中的落点", "解决什么问题"],
    [
        ["策略", "每种促销玩法封装为独立策略类，实现统一接口", "玩法可替换、可新增，互不干扰"],
        ["装饰", "多个促销按顺序逐层包裹价格计算", "玩法可叠加，且叠加顺序可控"],
        ["观察者", "交易完成事件发布订阅（结合阻塞队列）", "下单与后续处理解耦"],
        ["工厂", "促销策略的创建", "创建逻辑集中管理"],
        ["单例", "配置加载、连接池", "全局唯一实例"],
        ["状态", "订单与售后状态流转", "状态转换规则清晰"],
    ],
    widths=[2.2, 8.4, 6.0],
)
p("模式数量不在多，讲清与业务紧密贴合的若干模式，远胜于罗列用不上的模式。促销引擎中策略、装饰、反射三者协同，"
  "是本项目最核心的设计模式实践。")

h2("2.6 并发与互斥")
p("多用户同时抢购限量商品时，若扣减库存采用「先查询、再判断、后更新」三步式实现，两个线程之间会形成竞态条件，"
  "导致超卖。本项目的方案是数据库条件更新——把判断与扣减合并为一条原子 SQL：")
code("UPDATE t_product SET stock = stock - ? WHERE id = ? AND stock >= ?")
p("执行后检查影响行数：返回 1 表示扣减成功；返回 0 表示库存不足，回滚事务并返回失败。"
  "JVM 层按商品编号分段加锁是性能优化而非正确性保证——它只在单进程内有效，一旦服务端多实例部署就会失效，"
  "而数据库条件更新依然有效。本项目还计划以「超卖复现实验」作为并发正确性的直接证据。")

h2("2.7 同步与异步通信")
p("同步通信是请求—应答模式：客户端发出请求后阻塞等待应答，实现简单、结果直观。异步通信则允许请求发出后不等待，"
  "结果通过回调或事件返回。本项目采用「同步为主 + 推送通道」：核心功能以同步方式交付，保证进度；"
  "服务端在订单状态变更等事件发生时，通过长连接上的推送通道主动通知客户端，作为异步能力的落地点。")

# ================================================================ 第三章
h1("三、前端和后端概要设计")

h2("3.1 模块划分")
p("客户端（前端）模块：")
table(
    ["模块", "职责"],
    [
        ["ui 视图层", "登录/注册、商品列表、商品详情、购物车、结算、交易结果、售后、设置、店主端界面"],
        ["controller 控制器层", "各界面的事件调度，不写业务规则"],
        ["service 客户端业务层", "购物车副本维护、金额本地预演算"],
        ["comm 通信层", "协议编解码、连接管理"],
        ["dto 传输对象", "请求与应答的数据载体"],
    ],
    widths=[4.0, 12.6],
)
p("服务端（后端）模块：")
table(
    ["模块", "职责"],
    [
        ["comm 通信层", "Socket 监听、线程池、协议编解码、功能码分发"],
        ["service 业务层", "账号、商品、购物车、订单、促销、售后、事件发布"],
        ["dao 持久层", "DAO 接口 + JDBC 实现"],
        ["model / dto", "领域对象与传输对象"],
        ["util 工具层", "配置加载、日志、连接池"],
    ],
    widths=[4.0, 12.6],
)

h2("3.2 UI 设计")
p("客户端界面清单如下，其中店主端界面仅店主角色可见：")
table(
    ["界面", "主要内容", "主要交互"],
    [
        ["登录 / 注册", "账号、密码输入框", "校验通过后按角色进入主界面"],
        ["商品列表", "商品表格、搜索框、分页控件", "搜索、翻页、点击进入详情"],
        ["商品详情", "完整描述、促销活动", "选择数量、加入购物车"],
        ["购物车", "条目列表、复选框、数量增减", "修改后实时刷新小计"],
        ["结算", "结算单、优惠明细逐行展示", "确认提交订单"],
        ["交易结果", "成功/失败提示、订单号、计算明细", "查看明细、发起售后"],
        ["售后申请", "订单选择、退货原因表单", "提交申请"],
        ["设置", "界面主题选择", "即时换肤"],
        ["店主端界面", "商品维护、促销配置、订单查看、售后审核", "新增、修改、上/下架、审核"],
    ],
    widths=[3.0, 8.2, 5.4],
)
p("界面交互遵循三条原则：主界面按角色动态挂载菜单；客户端显示的预估金额一律带「预估」字样；"
  "界面提供 3 套主题可即时切换，主题切换仅影响表现层。")

h2("3.3 主要类的划分")
p("服务端主要类（按层划分）：")
table(
    ["层", "主要类"],
    [
        ["model 领域对象", "User、Product、CartItem、Order、OrderItem、Promotion、Coupon、AfterSale"],
        ["service 业务类", "UserService、ProductService、CartService、OrderService、PromotionService、AfterSaleService"],
        ["promotion 促销引擎", "PriceCalculator（接口）、DiscountStrategy、MarkdownStrategy、CouponStrategy、"
         "PromotionDecorator（抽象）、FullReduceDecorator、CouponDecorator、PromotionFactory"],
        ["dao 持久接口", "UserDao、ProductDao、CartDao、OrderDao、PromotionDao、CouponDao、AfterSaleDao（接口）"
         "＋ jdbc 包下的 JdbcXxxDao 实现"],
        ["comm 通信类", "ProtocolCodec、RequestHandler、ConnectionManager、ClientConnection"],
        ["event 事件类", "EventBus、OrderCompletedEvent、PromotionStartedEvent"],
        ["util 工具类", "Config、Logger、DBConnectionPool"],
    ],
    widths=[3.6, 13.0],
)
p("客户端主要类：")
table(
    ["层", "主要类"],
    [
        ["ui 视图", "LoginView、ProductListView、ProductDetailView、CartView、CheckoutView、ResultView、"
         "AfterSaleView、SettingView、OwnerView"],
        ["controller 控制器", "LoginController、ProductController、CartController、CheckoutController 等"],
        ["service 客户端业务", "CartModel（购物车副本）、PriceEstimator（金额预估）"],
        ["comm 通信", "ClientConnection、ProtocolCodec"],
        ["dto 传输对象", "各请求/应答对象"],
    ],
    widths=[3.6, 13.0],
)
p("设计要点：促销引擎是面向对象设计的核心——统一接口 PriceCalculator 下挂多个策略实现，装饰器逐层包裹，"
  "PromotionFactory 依据配置反射装配；领域对象与传输对象（dto）相互隔离，业务逻辑不得直接操作传输对象。")

h2("3.4 数据库设计")
p("本系统共 10 张表。字段类型统一使用 INTEGER / VARCHAR / TEXT / DECIMAL 表达；"
  "主键自增语法在 SQLite 与 MySQL 中不同，由两套建表脚本分别处理。各表字段如下：")

p("t_user　用户表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["username", "VARCHAR", "账号，唯一"],
    ["password_hash", "VARCHAR", "密码哈希"],
    ["salt", "VARCHAR", "随机盐值"],
    ["nickname", "VARCHAR", "昵称"],
    ["role", "VARCHAR", "角色：顾客 / 店主"],
    ["created_at", "TEXT", "创建时间（ISO-8601）"],
], widths=[3.2, 3.4, 10.0])

p("t_product　商品表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["name", "VARCHAR", "商品名称"],
    ["price", "DECIMAL", "单价"],
    ["stock", "INTEGER", "库存数量"],
    ["description", "TEXT", "描述"],
    ["status", "VARCHAR", "上架 / 下架"],
    ["created_at", "TEXT", "创建时间"],
], widths=[3.2, 3.4, 10.0])

p("t_cart_item　购物车条目表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["user_id", "INTEGER", "外键 → t_user"],
    ["product_id", "INTEGER", "外键 → t_product"],
    ["quantity", "INTEGER", "数量"],
    ["created_at", "TEXT", "加入时间"],
], widths=[3.2, 3.4, 10.0])

p("t_order　订单表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["order_no", "VARCHAR", "订单号，唯一"],
    ["user_id", "INTEGER", "外键 → t_user"],
    ["total_amount", "DECIMAL", "商品总额"],
    ["discount_amount", "DECIMAL", "优惠总额"],
    ["payable_amount", "DECIMAL", "应付金额"],
    ["status", "VARCHAR", "待付款/已付款/已完成/已取消/售后中/已退款"],
    ["created_at", "TEXT", "下单时间"],
    ["paid_at", "TEXT", "付款时间"],
], widths=[3.2, 3.4, 10.0])

p("t_order_item　订单明细表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["order_id", "INTEGER", "外键 → t_order"],
    ["product_id", "INTEGER", "外键 → t_product"],
    ["product_name", "VARCHAR", "商品名快照"],
    ["price_snapshot", "DECIMAL", "成交单价快照"],
    ["quantity", "INTEGER", "数量"],
    ["subtotal", "DECIMAL", "小计"],
], widths=[3.2, 3.4, 10.0])

p("t_promotion　促销活动表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["name", "VARCHAR", "活动名称"],
    ["type", "VARCHAR", "玩法标识（对应实现类名映射）"],
    ["params", "TEXT", "参数（JSON）"],
    ["status", "VARCHAR", "启用 / 停用"],
    ["created_at", "TEXT", "创建时间"],
], widths=[3.2, 3.4, 10.0])

p("t_promotion_scope　促销适用范围表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["promotion_id", "INTEGER", "外键 → t_promotion"],
    ["product_id", "INTEGER", "外键 → t_product"],
], widths=[3.2, 3.4, 10.0])

p("t_coupon　优惠券模板表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["name", "VARCHAR", "券名称"],
    ["type", "VARCHAR", "满减 / 抵扣"],
    ["threshold", "DECIMAL", "使用门槛"],
    ["amount", "DECIMAL", "面额"],
    ["status", "VARCHAR", "启用 / 停用"],
], widths=[3.2, 3.4, 10.0])

p("t_user_coupon　用户持有券表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["user_id", "INTEGER", "外键 → t_user"],
    ["coupon_id", "INTEGER", "外键 → t_coupon"],
    ["status", "VARCHAR", "未使用 / 已使用 / 已过期"],
    ["received_at", "TEXT", "领取时间"],
    ["used_at", "TEXT", "使用时间"],
], widths=[3.2, 3.4, 10.0])

p("t_after_sale　售后单表", indent=False, bold=True)
table(["字段", "类型", "说明"], [
    ["id", "INTEGER", "主键，自增"],
    ["order_id", "INTEGER", "外键 → t_order"],
    ["user_id", "INTEGER", "外键 → t_user"],
    ["reason", "TEXT", "申请原因"],
    ["status", "VARCHAR", "待审核 / 已通过 / 已拒绝"],
    ["created_at", "TEXT", "申请时间"],
    ["handled_at", "TEXT", "处理时间"],
], widths=[3.2, 3.4, 10.0])

p("表间关系：用户一对多购物车条目与订单；订单一对多订单明细；促销与商品多对多（经 t_promotion_scope）；"
  "用户与优惠券多对多（经 t_user_coupon）；订单与售后单一对一。")

h2("3.5 通信协议设计")
p("报文由 16 字节定长报文头与变长数据体组成：")
table(
    ["字段", "长度", "说明"],
    [
        ["魔数", "2 字节", "固定值 0xAA55，用于快速识别非法数据"],
        ["版本号", "1 字节", "协议版本，当前 0x01"],
        ["类型", "1 字节", "1=请求，2=应答，9=服务端推送"],
        ["功能码", "4 字节", "标识具体操作"],
        ["序列号", "4 字节", "请求与应答的配对标识"],
        ["数据长度", "4 字节", "数据体字节数，解决粘包与半包问题"],
        ["数据体", "变长", "JSON 文本序列化"],
    ],
    widths=[2.4, 2.0, 12.2],
)
p("主要功能码：")
table(
    ["功能码", "含义", "功能码", "含义"],
    [
        ["1001", "请求商品列表", "1005", "发起售后申请"],
        ["1002", "将商品加入购物车", "1006", "领取优惠券"],
        ["1003", "查询购物车", "1007", "用户登录"],
        ["1004", "提交结算", "1008", "查询界面主题配置"],
    ],
    widths=[2.4, 6.0, 2.4, 5.8],
)
p("编解码要点：按长度字段切分，不假设一次 read 即一条完整报文；长度以 UTF-8 字节数计；"
  "应答通过序列号配对，支持长连接上的并发请求；业务异常以错误码加消息在应答中传递，不断开连接。")

h2("3.6 主要逻辑流程图")
figure("图 1", "图 1　登录 / 注册流程图")
figure("图 2", "图 2　结算下单与防超卖流程图")
figure("图 3", "图 3　协议收发编解码流程图")

# ================================================================ 第四章
h1("四、其他")

h2("4.1 进度回顾")
p("对照第一篇制定的六周里程碑，当前进度如下：")
table(
    ["周次", "计划目标", "当前状态"],
    [
        ["第 1 周", "需求定稿与环境预热", "已完成：需求定稿，开发环境就绪，Socket 收发与 JDBC 读写各跑通一个最小例程"],
        ["第 2 周", "通信层打通", "进行中：【待我确认：第 2 周实际完成情况，如协议编解码、线程池、登录与商品列表】"],
        ["第 3 周", "数据层打通", "未开始"],
        ["第 4 周", "促销引擎", "未开始"],
        ["第 5 周", "难点攻坚", "未开始"],
        ["第 6 周", "收尾交付", "未开始"],
    ],
    widths=[2.0, 4.6, 10.0],
)

h2("4.2 遗留问题与后续计划")
p("遗留问题：指导教师姓名待补；3 套界面主题的具体配色与控件风格待实现阶段确定；六周里程碑的具体起止日期"
  "以课程安排为准。")
p("后续计划：第 3 周完成建库建表、DAO 接口与 JDBC 实现；第 4 周完成促销引擎（策略、装饰、反射协同）"
  "与结算流程；第 5 周攻坚异步事件分发与售后状态流转；第 6 周完成主题切换、测试补齐与演示材料。")

h2("4.3 个人收获与反思")
p("两周以来，通过需求分析与本篇设计，笔者对 C/S 架构、自定义通信协议、并发互斥与设计模式有了从概念到方案的"
  "落地理解。最大的收获是认识到「底层原理优先于框架」——先手写 JDBC 与自定义协议，才能真正理解映射与粘包半包，"
  "框架只是这些原理的封装。")
p("反思在于，最初对多店、真实支付等功能有求全的冲动，经范围管理明确「不做清单」后才聚焦到真正的考点上。"
  "这提醒笔者：边界清晰的中等规模系统，其工程价值高于边界模糊、处处半成品的大系统。")

doc.save(OUT)
print("saved:", OUT)
